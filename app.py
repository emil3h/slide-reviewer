"""
app.py - Streamlit interface for the Slide Reviewer (with per-issue checkboxes).

Rename your checker to exactly slide_check.py, put it next to this file, then:
    python -m streamlit run app.py
"""
import io
import zipfile
import tempfile
from pathlib import Path

import streamlit as st
from pptx import Presentation

import slide_check as sc

st.set_page_config(page_title="Slide Reviewer", page_icon="\U0001F4D0", layout="centered")


def review_path(path, marking_variant=None):
    prs = Presentation(str(path))
    return prs, sc.review(prs, marking_variant=marking_variant)


def fix_all_bytes(path):
    prs = Presentation(str(path))
    issues = sc.review(prs)
    sc.apply_selected(prs, issues, [i.id for i in issues if i.fixable])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


st.title("\U0001F4D0 Slide Reviewer")
st.caption("Checks each deck against the template rules (headings, body text, page "
           "numbers, CUI marking) and lets you approve fixes one by one.")

mode = st.radio("Mode", ["Single deck (upload)", "Batch (local folder)"], horizontal=True)

# ============================ SINGLE DECK ===================================
if mode == "Single deck (upload)":
    uploaded = st.file_uploader("Upload a presentation (.pptx)", type="pptx")
    if uploaded is None:
        st.info("Upload a .pptx to begin.")
        st.stop()

    # persist the upload to a temp file so it survives Streamlit reruns
    if st.session_state.get("deck_name") != uploaded.name:
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
        tmp.write(uploaded.getbuffer())
        tmp.flush()
        st.session_state["deck_path"] = tmp.name
        st.session_state["deck_name"] = uploaded.name
        st.session_state["marking_choice"] = "auto"  # reset override for the new deck

    variant_labels = {v["id"]: v["label"] for v in sc.CUI_VARIANTS}
    detected = sc.detect_marking_variant(Presentation(st.session_state["deck_path"]))
    choice = st.selectbox(
        "This deck's approved marking",
        ["auto"] + list(variant_labels),
        format_func=lambda vid: f"Auto-detect (currently: {variant_labels[detected]})"
                                 if vid == "auto" else variant_labels[vid],
        key="marking_choice",
        help="The deck must use exactly one marking variant throughout. Pick it once "
             "here and every slide is checked against it.")
    marking_variant = None if choice == "auto" else choice

    prs, issues = review_path(st.session_state["deck_path"], marking_variant)
    n_slides = len(list(prs.slides))
    st.subheader(f"{uploaded.name} \u2014 {n_slides} slides")

    if not issues:
        st.success("\u2705 No issues found. This deck passes all checks.")
        st.stop()

    fixable_ids = [i.id for i in issues if i.fixable]
    for iid in fixable_ids:
        st.session_state.setdefault(iid, True)

    n_fixable = len(fixable_ids)
    n_manual = len(issues) - n_fixable
    st.write(f"**{len(issues)} issue(s)** \u2014 {n_fixable} auto-fixable"
             + (f", {n_manual} manual" if n_manual else "") + ". "
             "Untick anything you don't want changed.")

    c1, c2 = st.columns(2)
    if c1.button("Select all"):
        for iid in fixable_ids:
            st.session_state[iid] = True
    if c2.button("Clear all"):
        for iid in fixable_ids:
            st.session_state[iid] = False

    st.divider()

    # render checkboxes grouped by slide
    last = None
    for i in issues:
        if i.slide != last:
            st.markdown(f"**Slide {i.slide}**")
            last = i.slide
        if i.fixable:
            st.checkbox(i.message, key=i.id)
        else:
            st.checkbox(f"{i.message}  *(fix manually in PowerPoint)*",
                        value=False, disabled=True, key=i.id)

    st.divider()
    selected = [i.id for i in issues if i.fixable and st.session_state.get(i.id)]
    if st.button(f"Apply {len(selected)} selected fix(es)", type="primary",
                 disabled=not selected):
        sc.apply_selected(prs, issues, selected)
        buf = io.BytesIO()
        prs.save(buf)
        st.success(f"Applied {len(selected)} fix(es).")
        st.download_button("\u2b07\ufe0f Download corrected deck", data=buf.getvalue(),
                           file_name=f"corrected_{uploaded.name}",
                           mime="application/vnd.openxmlformats-officedocument."
                                "presentationml.presentation")

# ============================ BATCH FOLDER ==================================
else:
    folder = st.text_input("Folder containing .pptx decks", value="submitted_decks")
    if not folder:
        st.stop()
    path = Path(folder)
    if not path.exists():
        st.error(f"Folder not found: {path.resolve()}")
        st.stop()
    decks = sorted(path.glob("*.pptx"))
    if not decks:
        st.warning("No .pptx files found there.")
        st.stop()

    st.write(f"Found **{len(decks)}** deck(s).")
    if st.button("Review all decks", type="primary"):
        rows = []
        progress = st.progress(0.0)
        for i, deck in enumerate(decks, start=1):
            _, issues = review_path(deck)
            rows.append({"Deck": deck.name, "Issues": len(issues)})
            progress.progress(i / len(decks))
        st.session_state["batch_rows"] = rows
        st.session_state["batch_folder"] = str(path)

    if "batch_rows" in st.session_state:
        rows = st.session_state["batch_rows"]
        total = sum(r["Issues"] for r in rows)
        st.subheader(f"Summary \u2014 {total} issue(s) across {len(rows)} decks")
        st.dataframe(rows, use_container_width=True, hide_index=True)
        st.divider()
        if st.button("Apply all fixes to every deck & download zip"):
            zbuf = io.BytesIO()
            fp = Path(st.session_state["batch_folder"])
            with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_DEFLATED) as zf:
                for deck in sorted(fp.glob("*.pptx")):
                    zf.writestr(f"corrected_{deck.name}", fix_all_bytes(deck))
            st.success("All decks corrected.")
            st.download_button("\u2b07\ufe0f Download all corrected decks (zip)",
                               data=zbuf.getvalue(), file_name="corrected_decks.zip",
                               mime="application/zip")
