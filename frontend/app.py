"""
app.py - Streamlit interface for the Slide Reviewer (with per-issue checkboxes).

Run from the repo root:
    python3 -m streamlit run frontend/app.py
"""
import io
import sys
import tempfile
from datetime import datetime
from itertools import groupby
from pathlib import Path

import streamlit as st
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import slide_check as sc

LOGO_PATH = Path(__file__).resolve().parent / "logo.png"
LOGO_IMAGE = str(LOGO_PATH) if LOGO_PATH.exists() else None

st.set_page_config(page_title="Slide Reviewer",
                   page_icon=LOGO_IMAGE if LOGO_IMAGE is not None else None,
                   layout="centered")


def inject_custom_styles():
    """Purely cosmetic: card elevation, tighter alert/metric chrome, and
    hover/transition polish on the buttons tagged with `key=` below. `.streamlit/
    config.toml` owns the actual color/font theme; this only fills in what
    config.toml can't reach."""
    st.html("""
    <style>
        .block-container {
            padding-top: 2rem !important;
            padding-bottom: 3rem !important;
            max-width: 900px !important;
        }

        /* st.container(border=True) -> elevated card */
        div[data-testid="stVerticalBlockBorderWrapper"] {
            border-radius: 14px !important;
            box-shadow: 0 1px 3px rgba(15, 23, 42, 0.06), 0 1px 2px rgba(15, 23, 42, 0.04) !important;
            transition: box-shadow 0.2s ease !important;
        }
        div[data-testid="stVerticalBlockBorderWrapper"]:hover {
            box-shadow: 0 6px 16px rgba(15, 23, 42, 0.08) !important;
        }

        div[data-testid="stAlertContainer"] {
            border-radius: 10px !important;
        }

        div[data-testid="stMetric"] {
            background: rgba(37, 99, 235, 0.05);
            border-radius: 10px;
            padding: 0.75rem 1rem;
        }

        /* Primary CTA */
        .st-key-apply-fixes-btn button {
            background-color: #2563EB !important;
            color: white !important;
            border-radius: 8px !important;
            border: none !important;
            font-weight: 600 !important;
            transition: all 0.15s ease !important;
        }
        .st-key-apply-fixes-btn button:hover {
            background-color: #1D4ED8 !important;
            transform: translateY(-1px);
            box-shadow: 0 4px 10px rgba(37, 99, 235, 0.35) !important;
        }

        /* Toolbar buttons */
        .st-key-select-all-btn button,
        .st-key-clear-all-btn button {
            border-radius: 8px !important;
            transition: all 0.15s ease !important;
        }
    </style>
    """)


inject_custom_styles()

with st.sidebar:
    st.header("Review Rubric")
    st.markdown(
        """
        **Font Rules**
        - **Heading:** Arial, 24 pt, bold, black
        - **Body:** Arial, 16 pt or larger

        **Page Number**
        - Bottom-right, Arial, 8 pt
        - Sequential order

        **Logo**
        - Top-left
        - Width of at least 3 inches

        **Markings**
        - Exactly one variant
        - Exact wording and Arial formatting
        """
    )

def review_path(path, marking_variant=None):
    prs = Presentation(str(path))
    return prs, sc.review(prs, marking_variant=marking_variant)

def build_deck_context(deck_name, n_slides, issues, applied_ids=None):
    """Turn the actual review results into text the chat model can reason
    over: every issue, which slide it's on, whether it's auto-fixable, and
    (once the user has clicked Apply) which ones were actually applied."""
    applied_ids = set(applied_ids or [])
    lines = [
        f"Deck name: {deck_name}",
        f"Total slides: {n_slides}",
        f"Total issues found: {len(issues)}",
        "",
        "Full list of issues (grouped by slide):",
    ]
    if not issues:
        lines.append("(none — deck passes all checks)")
    else:
        last_slide = None
        for i in issues:
            if i.slide != last_slide:
                lines.append(f"\nSlide {i.slide}:")
                last_slide = i.slide
            status = "auto-fixable" if i.fixable else "requires manual fix"
            if i.id in applied_ids:
                status = "auto-fixable — FIX APPLIED by user"
            lines.append(f"  - [{i.category}] {i.message} ({status})")

    return "\n".join(lines)


def get_chathpc_response(messages, context=""):
    """
    Calls the ChatHPC API using credentials stored in st.secrets.
    Includes truststore injection to handle JPL self-signed certificates.

    `messages` is the full running chat history (list of {"role","content"}
    dicts) so the model has both the deck context and prior turns, not just
    the latest question in isolation.
    """
    import requests
    try:
        import truststore
        truststore.inject_into_ssl()
    except ImportError:
        # If truststore isn't installed, we'll try to proceed, but might hit SSL errors
        pass
    
    api_key = st.secrets.get("CHATHPC_API_KEY")
    endpoint = st.secrets.get("CHATHPC_API_ENDPOINT")
    
    if not api_key or not endpoint:
        return "Error: ChatHPC API credentials not configured in `.streamlit/secrets.toml`."

    system_prompt = (
        "You are a helpful assistant embedded in a PowerPoint compliance review tool. "
        "Answer questions about the specific deck the user uploaded, using the issue "
        "list below as ground truth. Don't invent issues that aren't listed, and be "
        "specific about slide numbers when relevant.\n\n"
        f"Context:\n{context}"
    )

    try:
        payload = {
            "model": "gemma4:31b-128k",
            "messages": [{"role": "system", "content": system_prompt}] + messages,
            "temperature": 0.7,
            "chat_template_kwargs": {"enable_thinking": True}
        }
        headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}
        
        response = requests.post(endpoint, json=payload, headers=headers, timeout=60)
        response.raise_for_status()
        return response.json().get("choices", [{}])[0].get("message", {}).get("content", "No response content found.")
    except Exception as e:
        return f"ChatHPC Error: {str(e)}"


def build_report(deck_name, n_slides, issues, applied_ids):
    """Plain-text summary of what was auto-fixed vs. what still needs a human
    to go into PowerPoint, grouped by slide."""
    applied_ids = set(applied_ids)
    implemented = [i for i in issues if i.id in applied_ids]
    manual = [i for i in issues if i.id not in applied_ids]

    lines = [
        "SLIDE REVIEWER — CHANGE REPORT",
        f"Deck: {deck_name} ({n_slides} slides)",
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        "",
    ]

    def section(title, items, note_unselected=False):
        lines.append(f"{title} ({len(items)})")
        lines.append("=" * (len(title) + len(str(len(items))) + 3))
        if not items:
            lines.append("(none)")
        else:
            last = None
            for i in items:
                if i.slide != last:
                    lines.append(f"\nSlide {i.slide}:")
                    last = i.slide
                tag = ""
                if note_unselected and i.fixable:
                    tag = "  [auto-fixable, but was left unselected]"
                lines.append(f"  - {i.message}{tag}")
        lines.append("")

    section("IMPLEMENTED FIXES", implemented)
    section("NEEDS MANUAL ACTION", manual, note_unselected=True)

    return "\n".join(lines)




header_logo, header_title = st.columns([1, 4], vertical_alignment="center")
if LOGO_IMAGE is not None:
    header_logo.image(LOGO_IMAGE, width=140)
header_title.title("Slide Reviewer")
st.caption("Checks each deck against the template rules (headings, body text, page "
           "numbers, markings) and lets you approve fixes one by one.")

# ============================ SINGLE DECK ===================================
with st.container(border=True):
    st.subheader("Upload deck")
    uploaded = st.file_uploader("Upload a presentation (.pptx)", type="pptx")
    if uploaded is None:
        st.info("Upload a .pptx to begin.", icon=":material/upload_file:")
        st.stop()

# persist the upload to a temp file so it survives Streamlit reruns
if st.session_state.get("deck_name") != uploaded.name:
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
    tmp.write(uploaded.getbuffer())
    tmp.flush()
    st.session_state["deck_path"] = tmp.name
    st.session_state["deck_name"] = uploaded.name
    # Auto-detect marking for the new deck
    detected = sc.detect_marking_variant(Presentation(st.session_state["deck_path"]))
    st.session_state["marking_choice"] = detected
elif "marking_choice" not in st.session_state:
    detected = sc.detect_marking_variant(Presentation(st.session_state["deck_path"]))
    st.session_state["marking_choice"] = detected

marking_variant = st.session_state["marking_choice"]
with st.spinner("Analyzing deck...", show_time=True):
    prs, issues = review_path(st.session_state["deck_path"], marking_variant)
n_slides = len(list(prs.slides))
st.subheader(f"{uploaded.name} — {n_slides} slides")

if not issues:
    st.success("No issues found. This deck passes all checks.", icon=":material/check_circle:")
    st.stop()

fixable_ids = [i.id for i in issues if i.fixable]
for iid in fixable_ids:
    st.session_state.setdefault(iid, True)

n_fixable = len(fixable_ids)
n_manual = len(issues) - n_fixable

with st.container(border=True):
    st.subheader("Overview")
    m1, m2, m3 = st.columns(3)
    m1.metric("Issues found", len(issues))
    m2.metric("Auto-fixable", n_fixable)
    m3.metric("Manual", n_manual)
    st.caption("Untick anything below you don't want changed.")

    c1, c2, c3 = st.columns([1, 1, 2])
    if c1.button("Select all", key="select-all-btn", use_container_width=True):
        for iid in fixable_ids:
            st.session_state[iid] = True
        st.toast("All fixable issues selected")
    if c2.button("Clear all", key="clear-all-btn", use_container_width=True):
        for iid in fixable_ids:
            st.session_state[iid] = False
        st.toast("All selections cleared")

    hide_compliant = c3.checkbox("Hide compliant slides", value=False)

st.divider()

# render checkboxes grouped by slide, one card per slide
for slide_no, slide_issues in groupby(issues, key=lambda i: i.slide):
    slide_issues = list(slide_issues)

    if hide_compliant and not slide_issues:
        continue

    with st.container(border=True):
        st.markdown(f"**Slide {slide_no}**")
        for i in slide_issues:
            if i.fixable:
                st.checkbox(i.message, key=i.id)
            else:
                st.warning(f"**Manual Fix Required:** {i.message}", icon=":material/build:")

st.divider()
selected = [i.id for i in issues if i.fixable and st.session_state.get(i.id)]
custom_labels = {v["id"]: v["label"] for v in sc.CUI_VARIANTS}
# Override with user-requested labels
custom_labels["not_cui_footer"] = "Reviewed and determined not to contain CUI"
custom_labels["cui_header"] = "CUI header banner"

st.selectbox(
    "Approved marking for this deck",
    options=list(custom_labels.keys()),
    format_func=lambda x: custom_labels[x],
    key="marking_choice",
    help="The deck must use exactly one marking variant throughout. Changing this will re-run the review.",
)

if st.button(f"Apply {len(selected)} selected fix(es)", key="apply-fixes-btn", type="primary",
             disabled=not selected, use_container_width=True):
    with st.spinner("Applying fixes..."):
        sc.apply_selected(prs, issues, selected)
        buf = io.BytesIO()
        prs.save(buf)
    st.session_state["last_applied_ids"] = selected
    st.toast(f"Applied {len(selected)} fix(es)", icon=":material/check_circle:")
    st.success(f"Applied {len(selected)} fix(es).", icon=":material/check_circle:")

    dl1, dl2 = st.columns(2)
    dl1.download_button("Download corrected deck", data=buf.getvalue(),
                       file_name=f"corrected_{uploaded.name}",
                       mime="application/vnd.openxmlformats-officedocument."
                             "presentationml.presentation",
                       use_container_width=True)
    report_text = build_report(uploaded.name, n_slides, issues, selected)
    dl2.download_button("Download change report", data=report_text,
                       file_name=f"report_{Path(uploaded.name).stem}.txt",
                       mime="text/plain", use_container_width=True)

st.divider()

# ============================ CHATHPC AGENT ===================================
st.subheader("Ask ChatHPC")
st.info("Ask questions about the deck review or general slide formatting guidelines.", icon=":material/forum:")

if "chat_history" not in st.session_state:
    st.session_state.chat_history = []

# Display chat history
for message in st.session_state.chat_history:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])

# Chat input
if prompt := st.chat_input("How can I improve my slides?"):
    # Display user message
    with st.chat_message("user"):
        st.markdown(prompt)
    st.session_state.chat_history.append({"role": "user", "content": prompt})

    # Get response from ChatHPC
    with st.chat_message("assistant"):
        with st.spinner("ChatHPC is thinking..."):
            # Full issue list (+ whatever fixes have been applied so far this
            # session), plus the running conversation, so the model actually
            # knows what was found and can answer follow-up questions.
            context = build_deck_context(
                uploaded.name, n_slides, issues,
                applied_ids=st.session_state.get("last_applied_ids", []),
            )
            response = get_chathpc_response(st.session_state.chat_history, context=context)
            st.markdown(response)
    st.session_state.chat_history.append({"role": "assistant", "content": response})