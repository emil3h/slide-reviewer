"""
slide_check.py - formatting checker with per-issue, individually-applyable fixes.

Rules (applied to every slide except the title slide, and skipping "Back-up Slides"
dividers for font checks):
  1. Page number present (typed number or slide-number placeholder), numbered in order.
  2. CUI marking present, worded exactly, formatted Arial 8 bold no highlight.
  3. Heading (top-most text box) Arial 24.
  4. All other body text Arial 16.

Public API used by app.py:
  review(prs) -> list[Issue]           # each Issue has .id, .slide, .message, .fixable
  apply_selected(prs, issues, ids)     # apply only the issues whose id is in ids

CLI:
  python slide_check.py deck.pptx            # report
  python slide_check.py deck.pptx --fix       # apply ALL fixable issues, write a copy
"""
import argparse
import re
import sys
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import PP_PLACEHOLDER

# ---- rules -----------------------------------------------------------------
HEADING_FONT, HEADING_SIZE = "Arial", 24.0
BODY_FONT, BODY_SIZE = "Arial", 16.0
CUI_TEXT = "Reviewed and determined not to contain CUI"
CUI_FONT, CUI_SIZE = "Arial", 8.0
SKIP_FIRST_SLIDE = True
TRUST_INHERITED = True
MASTER_PROVIDES_PAGE_NUMBERS = False


@dataclass
class Issue:
    id: str
    slide: int
    category: str
    message: str
    fixable: bool
    fix: dict = field(default_factory=dict)


# ---- text / shape helpers --------------------------------------------------
def runs(shape):
    if not shape.has_text_frame:
        return []
    out = []
    for p in shape.text_frame.paragraphs:
        out.extend(p.runs)
    return out


def text_of(shape):
    return shape.text_frame.text.strip() if shape.has_text_frame else ""


def norm(s):
    return re.sub(r"\s+", " ", s.lower()).strip().rstrip(".")


def is_pageno(shape):
    if not shape.has_text_frame or shape.left is None or shape.top is None:
        return False
    if not re.fullmatch(r"\d{1,3}", text_of(shape)):
        return False
    return Emu(shape.left).inches > 7.0 and Emu(shape.top).inches > 6.0


def is_slidenum(shape):
    try:
        if shape.is_placeholder and \
                shape.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            return True
    except Exception:
        pass
    try:
        return "slidenum" in shape._element.xml.lower()
    except Exception:
        return False


def is_number_text(shape):
    return shape.has_text_frame and bool(re.fullmatch(r"\d{1,3}", text_of(shape)))


def is_footer_ph(shape):
    try:
        return shape.is_placeholder and \
            shape.placeholder_format.type == PP_PLACEHOLDER.FOOTER
    except Exception:
        return False


def is_pageno_carrier(shape):
    return is_pageno(shape) or is_number_text(shape) or is_slidenum(shape) \
        or is_footer_ph(shape)


def _safe_xml(shape):
    try:
        return shape._element.xml.lower()
    except Exception:
        return ""


def has_pageno(slide):
    if MASTER_PROVIDES_PAGE_NUMBERS:
        return True
    for s in slide.shapes:
        if is_pageno(s) or is_number_text(s) or is_slidenum(s):
            return True
        if is_footer_ph(s) and ("slidenum" in _safe_xml(s)
                                or any(c.isdigit() for c in text_of(s))):
            return True
    return False


def pageno_value(slide):
    for s in slide.shapes:
        if s.has_text_frame and re.fullmatch(r"\d{1,3}", text_of(s)):
            return int(text_of(s))
    return None


def is_cui_exact(shape):
    return norm(text_of(shape)) == norm(CUI_TEXT)


def is_cui_like(shape):
    return bool(text_of(shape)) and \
        SequenceMatcher(None, norm(text_of(shape)), norm(CUI_TEXT)).ratio() > 0.7


def is_backup_divider(slide):
    texts = [text_of(s) for s in slide.shapes
             if s.has_text_frame and text_of(s)
             and not is_pageno_carrier(s) and not is_cui_like(s)]
    flat = re.sub(r"\s+", " ", " ".join(texts).lower().replace("-", " ")).strip()
    return flat in ("back up slides", "backup slides", "back up slide", "backup slide")


def _rPr(run):
    return run._r.find(qn("a:rPr"))


def has_highlight(run):
    rPr = _rPr(run)
    return rPr is not None and rPr.find(qn("a:highlight")) is not None


def remove_highlight(run):
    rPr = _rPr(run)
    if rPr is not None:
        h = rPr.find(qn("a:highlight"))
        if h is not None:
            rPr.remove(h)


def inherited_shapes(slide):
    out = []
    try:
        layout = slide.slide_layout
        out.extend(layout.shapes)
        out.extend(layout.slide_master.shapes)
    except Exception:
        pass
    return out


# ---- theme font-token resolution -------------------------------------------
_THEME_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_theme_cache = {}


def _theme_fonts(slide):
    try:
        master = slide.slide_layout.slide_master
    except Exception:
        return (None, None)
    key = id(master)
    if key in _theme_cache:
        return _theme_cache[key]
    major = minor = None
    try:
        from lxml import etree
        blob = master.part.part_related_by(_THEME_RT).blob
        root = etree.fromstring(blob)
        mj = root.find(f".//{_A_NS}fontScheme/{_A_NS}majorFont/{_A_NS}latin")
        mn = root.find(f".//{_A_NS}fontScheme/{_A_NS}minorFont/{_A_NS}latin")
        major = mj.get("typeface") if mj is not None else None
        minor = mn.get("typeface") if mn is not None else None
    except Exception:
        pass
    _theme_cache[key] = (major, minor)
    return (major, minor)


def effective_font_name(name, slide):
    if name in ("+mj-lt", "+mj-ea", "+mj-cs"):
        return _theme_fonts(slide)[0]
    if name in ("+mn-lt", "+mn-ea", "+mn-cs"):
        return _theme_fonts(slide)[1]
    return name


def heading_shape(slide):
    best, best_top = None, None
    for shp in slide.shapes:
        if not shp.has_text_frame or not text_of(shp):
            continue
        if is_pageno_carrier(shp) or is_cui_like(shp) or shp.top is None:
            continue
        top = Emu(shp.top).inches
        if best_top is None or top < best_top:
            best, best_top = shp, top
    return best


# ---- review (produces Issues) ----------------------------------------------
def _font_issues(shape, want_font, want_size, where, slide, slide_no, target, idp, issues):
    for r in runs(shape):
        name = effective_font_name(r.font.name, slide)
        if name is not None and name != want_font:
            issues.append(Issue(f"{idp}-font", slide_no, "font",
                                f"{where} font is {name}; should be {want_font}.", True,
                                {"op": "font_name", "value": want_font, **target}))
            break
    for r in runs(shape):
        if r.font.size is not None and round(r.font.size.pt, 1) != want_size:
            issues.append(Issue(f"{idp}-size", slide_no, "size",
                                f"{where} size is {round(r.font.size.pt,1)}pt; "
                                f"should be {int(want_size)}pt.", True,
                                {"op": "font_size", "value": want_size, **target}))
            break


def _review_slide(slide, slide_no, issues):
    # page number presence
    if not has_pageno(slide):
        issues.append(Issue(f"s{slide_no}-pageno-missing", slide_no, "page_number",
                            "No page number found on the slide.", True,
                            {"op": "add_pageno", "n": slide_no}))

    # CUI marking
    slide_exact = [s for s in slide.shapes if is_cui_exact(s)]
    slide_like = [s for s in slide.shapes if is_cui_like(s) and not is_cui_exact(s)]
    if slide_exact:
        sh = slide_exact[0]
        defects = []
        rs = runs(sh)
        for r in rs:
            n = effective_font_name(r.font.name, slide)
            if n is not None and n != CUI_FONT:
                defects.append(f"font {n}")
                break
        sizes = [round(r.font.size.pt, 1) for r in rs if r.font.size is not None]
        if not sizes or any(sz != CUI_SIZE for sz in sizes):
            defects.append("not 8pt")
        if not all(r.font.bold for r in rs):
            defects.append("not bold")
        if any(has_highlight(r) for r in rs):
            defects.append("has highlight")
        if defects:
            issues.append(Issue(f"s{slide_no}-cui-format", slide_no, "cui",
                                f"CUI marking formatting is off ({', '.join(defects)}); "
                                f"should be Arial 8, bold, no highlight.", True,
                                {"op": "fix_cui_format", "shape_id": sh.shape_id}))
    elif any(is_cui_exact(s) for s in inherited_shapes(slide)):
        pass
    elif slide_like:
        issues.append(Issue(f"s{slide_no}-cui-wording", slide_no, "cui",
                            f"CUI marking wording differs; must read exactly: "
                            f"\u201c{CUI_TEXT}\u201d.", True,
                            {"op": "fix_cui_text", "shape_id": slide_like[0].shape_id}))
    else:
        issues.append(Issue(f"s{slide_no}-cui-missing", slide_no, "cui",
                            f"Missing required CUI marking: \u201c{CUI_TEXT}\u201d.", True,
                            {"op": "add_cui"}))

    # section divider: skip heading/body font checks
    if is_backup_divider(slide):
        return

    head = heading_shape(slide)
    if head is None:
        issues.append(Issue(f"s{slide_no}-noheading", slide_no, "heading",
                            "No heading text box found.", False, {}))
    else:
        _font_issues(head, HEADING_FONT, HEADING_SIZE, "Heading", slide, slide_no,
                     {"target": "heading"}, f"s{slide_no}-heading", issues)

    for shp in slide.shapes:
        if not shp.has_text_frame or not text_of(shp):
            continue
        if is_pageno_carrier(shp) or is_cui_like(shp):
            continue
        if head is not None and shp.shape_id == head.shape_id:
            continue
        _font_issues(shp, BODY_FONT, BODY_SIZE, "Body text", slide, slide_no,
                     {"target": "shape", "shape_id": shp.shape_id},
                     f"s{slide_no}-body-{shp.shape_id}", issues)


def _review_order(slides, start, issues):
    anchor_pos = anchor_num = None
    for idx, slide in enumerate(slides, start=1):
        val = pageno_value(slide)
        if val is None:
            continue
        if anchor_num is None:
            anchor_pos, anchor_num = idx, val
            continue
        expected = anchor_num + (idx - anchor_pos)
        if val != expected and idx >= start:
            # find the typed number shape to renumber
            sid = None
            for s in slide.shapes:
                if s.has_text_frame and re.fullmatch(r"\d{1,3}", text_of(s)):
                    sid = s.shape_id
                    break
            issues.append(Issue(f"s{idx}-pageno-order", idx, "page_number",
                                f"Page number is {val} but should be {expected} "
                                f"for its position (out of order).", sid is not None,
                                {"op": "set_number", "shape_id": sid, "value": expected}))


def review(prs):
    slides = list(prs.slides)
    start = 2 if SKIP_FIRST_SLIDE else 1
    issues = []
    for idx, slide in enumerate(slides, start=1):
        if idx < start:
            continue
        _review_slide(slide, idx, issues)
    _review_order(slides, start, issues)
    issues.sort(key=lambda i: i.slide)
    return issues


# ---- fix primitives --------------------------------------------------------
def set_font(shape, name, size):
    for r in runs(shape):
        r.font.name = name
        r.font.size = Pt(size)


def set_text(shape, new_text):
    rs = runs(shape)
    if rs:
        rs[0].text = new_text
        for extra in rs[1:]:
            extra.text = ""
    elif shape.has_text_frame:
        shape.text_frame.paragraphs[0].add_run().text = new_text


def add_pageno(slide, n):
    box = slide.shapes.add_textbox(Inches(9.0), Inches(7.0), Inches(0.6), Inches(0.3))
    box.text_frame.margin_left = 0
    box.text_frame.margin_top = 0
    r = box.text_frame.paragraphs[0].add_run()
    r.text = str(n); r.font.name = "Arial"; r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)


def add_cui_box(slide):
    box = slide.shapes.add_textbox(Inches(3.5), Inches(7.15), Inches(6.3), Inches(0.3))
    box.text_frame.margin_top = 0
    box.text_frame.margin_bottom = 0
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.add_run().text = CUI_TEXT
    return box


def format_cui_run(r):
    r.font.name = CUI_FONT
    r.font.size = Pt(CUI_SIZE)
    r.font.bold = True
    remove_highlight(r)


def _shape_by_id(slide, sid):
    for s in slide.shapes:
        if s.shape_id == sid:
            return s
    return None


def _apply_fix(slide, fix):
    op = fix.get("op")
    if op == "font_name":
        sh = heading_shape(slide) if fix.get("target") == "heading" \
            else _shape_by_id(slide, fix.get("shape_id"))
        if sh:
            for r in runs(sh):
                r.font.name = fix["value"]
    elif op == "font_size":
        sh = heading_shape(slide) if fix.get("target") == "heading" \
            else _shape_by_id(slide, fix.get("shape_id"))
        if sh:
            for r in runs(sh):
                r.font.size = Pt(fix["value"])
    elif op == "add_pageno":
        add_pageno(slide, fix["n"])
    elif op == "set_number":
        sh = _shape_by_id(slide, fix.get("shape_id"))
        if sh:
            set_text(sh, str(fix["value"]))
    elif op == "add_cui":
        box = add_cui_box(slide)
        for r in runs(box):
            format_cui_run(r)
    elif op == "fix_cui_text":
        sh = _shape_by_id(slide, fix.get("shape_id"))
        if sh:
            set_text(sh, CUI_TEXT)
            for r in runs(sh):
                format_cui_run(r)
    elif op == "fix_cui_format":
        sh = _shape_by_id(slide, fix.get("shape_id"))
        if sh:
            for r in runs(sh):
                format_cui_run(r)


def apply_selected(prs, issues, ids):
    id_set = set(ids)
    slides = list(prs.slides)
    for i in issues:
        if i.id in id_set and i.fixable and 1 <= i.slide <= len(slides):
            _apply_fix(slides[i.slide - 1], i.fix)


# ---- CLI -------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--fix", action="store_true")
    args = ap.parse_args()
    path = Path(args.deck)
    if not path.exists():
        sys.exit(f"File not found: {path}")

    prs = Presentation(str(path))
    issues = review(prs)
    skipped = " (slide 1 skipped as the title slide)" if SKIP_FIRST_SLIDE else ""
    if not issues:
        print(f"No issues found. All checked slides pass{skipped}.")
    else:
        print(f"\nFound {len(issues)} issue(s) in {path.name}{skipped}:")
        print("-" * 62)
        last = None
        for i in issues:
            if i.slide != last:
                print(f"\nSlide {i.slide}")
                last = i.slide
            tag = "" if i.fixable else " (manual)"
            print(f"  - {i.message}{tag}")
        print("-" * 62)

    if args.fix:
        apply_selected(prs, issues, [i.id for i in issues if i.fixable])
        out = path.with_name(f"corrected_{path.name}")
        prs.save(str(out))
        print(f"\nApplied all fixable issues. Wrote: {out}")
    elif issues:
        print("\nRun again with --fix to write a corrected copy.")


if __name__ == "__main__":
    main()
